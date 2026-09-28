"""Full Modeler -> Validator -> remediation -> sign-off run with real code
execution; only the LLM is scripted."""
from __future__ import annotations

import io

import pytest
from docx import Document

from studio.library import document_text
from studio.orchestrator import InputError
from studio.schemas import ProjectInput, Signoff

pytestmark = pytest.mark.slow


def _concept_paper() -> bytes:
    doc = Document()
    doc.add_heading("Mortgage PD concept paper", 0)
    doc.add_heading("Test sample", 1)
    doc.add_paragraph("The PD model must be evaluated on an out-of-time test sample covering the most recent two years.")
    doc.add_heading("Drivers", 1)
    doc.add_paragraph("Affordability (income) and leverage (debt ratio) must be considered as risk drivers.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_develop_and_validate_loop(studio, loans_csv):
    inp = ProjectInput(mode="develop_and_validate", title="Mortgage PD", brief="Build a 12-month PD model for retail mortgages.",
                       frameworks=["eu_banking"], max_rounds=3, challenger="auto")
    project = studio.create_project(inp, data=[("loans.csv", loans_csv)], concept_papers=[("concept.docx", _concept_paper())])
    project = studio.run(project.project_id)

    assert project.status == "awaiting_signoff", project.error
    assert len(project.rounds) == 2, "round 1 should fail and trigger one remediation round"
    r1, r2 = project.rounds

    # Round 1: weak model, missing planned test, review finding -> remediation required.
    ids1 = {f.finding_id for f in r1.validator.findings if f.status == "open"}
    assert "F-R03-MT-05" in ids1
    assert "F-R05-auc" in ids1 and "F-R04-MT-02" in ids1
    assert "F-G-credit-risk-gini" in ids1  # deterministic EBA gate on Gini
    assert any(fid.startswith("F-Q-") for fid in ids1)
    assert r1.validator.outcome == "remediation_required"
    assert r1.validator.replication and r1.validator.replication.reproduced
    # A fabricated citation is dropped; the attachment requirement keeps a real locator.
    att = [r for r in r1.validator.requirements if r.req_id.startswith("ATT-")]
    assert [r.req_id for r in att] == ["ATT-01"]
    assert att[0].citation and att[0].citation.locator.startswith("section")
    q = next(f for f in r1.validator.findings if f.finding_id.startswith("F-Q-"))
    assert q.citations == []

    # Round 2: modeler responded, validator closed the findings.
    assert r2.modeler.responses and all(r.action == "fixed" for r in r2.modeler.responses)
    closed = {f.finding_id for f in r2.validator.findings if f.status == "closed"}
    assert {"F-R03-MT-05", "F-R05-auc", "F-R04-MT-02"} <= closed
    assert all(f.closed_round == 2 for f in r2.validator.findings if f.finding_id in closed)
    assert r2.validator.outcome in ("approved", "approved_with_conditions")
    assert project.final_outcome == r2.validator.outcome

    # Challenger ran once (attachments present) and was reused with the new champion.
    assert r1.validator.challenger and r1.validator.challenger.pipeline_ok
    assert r2.validator.challenger.champion_metrics["auc"] > r1.validator.challenger.champion_metrics["auc"]
    assert studio.llm.calls.count("challenger") == 1
    assert studio.llm.calls.count("independent") == 1

    # Documents exist and carry executed results.
    pdir = studio.store.project_dir(project.project_id)
    text = document_text(pdir / r2.modeler.document_docx)
    assert "Response to validation findings" in text
    assert "MT-05" in text
    report = document_text(pdir / r2.validator.report_docx)
    assert "Requirement coverage" in report and "ATT-01" in report

    # Human sign-off is recorded and the report is regenerated.
    signed = studio.signoff(project.project_id, Signoff(name="Jane Validator", role="Head of validation"))
    assert signed.status == "signed_off"
    assert "Jane Validator" in document_text(pdir / signed.rounds[-1].validator.report_docx)
    with pytest.raises(InputError):
        studio.signoff(project.project_id, Signoff(name="jane validator"))


def test_inputs_are_checked(studio, loans_csv):
    inp = ProjectInput(mode="validate_external", title="x", brief="y")
    with pytest.raises(InputError):
        studio.create_project(inp, data=[("loans.csv", loans_csv)])
    with pytest.raises(InputError):
        studio.create_project(ProjectInput(mode="develop_and_validate", title="x", brief="y"),
                              data=[("loans.exe", b"x")])


def test_validate_external_package(studio, loans_csv):
    import zipfile

    from .scripted_llm import EXTERNAL_CODE

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("vendor_model/pdmodel/__init__.py", "")
        zf.writestr("vendor_model/pdmodel/core.py", EXTERNAL_CODE)
        zf.writestr("vendor_model/README.md", "# Vendor PD model\n\nLogistic regression. Out-of-time AUC is 0.80.\n")
    inp = ProjectInput(mode="validate_external", title="Vendor PD", brief="Validate the vendor PD model.",
                       frameworks=["eu_banking", "eu_ai_act"], challenger="off")
    project = studio.create_project(inp, data=[("loans.csv", loans_csv)], package=[("vendor.zip", buf.getvalue())])
    project = studio.run(project.project_id)
    assert project.status == "awaiting_signoff", project.error
    assert len(project.rounds) == 1
    rnd = project.rounds[0]
    rep = rnd.validator.replication
    assert rep.pipeline_ok and not rep.reproduced and rep.mismatches == ["auc"]  # documented 0.80 is overstated
    ids = {f.finding_id for f in rnd.validator.findings}
    assert "F-R02-auc" in ids
    assert any(r.req_id.startswith("AIA-") for r in rnd.validator.requirements)
    assert "external-intake" in studio.llm.calls and "implement" not in studio.llm.calls
