"""Golden/snapshot tests for the PPTX/DOCX report templates.

Renders a fixed fixture ValidationReport and diffs its extracted structured
content (slide/paragraph/table text -- not raw bytes, which would be
sensitive to timestamps/embedded-image encoding) against a checked-in
reference under tests/golden/, to catch unintended layout/content
regressions as tools/pptx_report.py and tools/docx_report.py evolve.

If a golden file doesn't exist yet, it's created from the current output and
the test is skipped with a note to re-run -- the standard snapshot-testing
bootstrap pattern. Delete a golden file and re-run to intentionally re-baseline
after a deliberate template change.
"""
import json
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from agents.risk_schemas import ValidationFinding, ValidationReport
from tools.docx_report import ValidationWordReportBuilder
from tools.pptx_report import ValidationDeckBuilder

GOLDEN_DIR = Path(__file__).parent / "golden"


def _fixture_report() -> ValidationReport:
    finding = ValidationFinding(
        finding_id="finding-0001",
        domain="credit_risk", area="Calibration and back-testing", verdict="non_compliant",
        severity="critical", description="PSI breach detected in the retail PD model case file.",
        recommendation="Recalibrate the model and reassess quarterly.", owner="Model Owner",
        remediation_deadline_days=30,
    )
    return ValidationReport(
        report_id="golden-report-0001",
        domain="credit_risk", title="Credit Risk Validation Report -- PD-RETAIL-01",
        scope="Scope text", methodology="Methodology text",
        entity_under_review="PD-RETAIL-01 (mortgages)", reporting_period="2026Q2",
        findings=[finding], quantitative_results={"psi": 0.30, "gini": 0.35},
        overall_rating="non_compliant", overall_conclusion="Critical PSI breach requires recalibration.",
        generated_at="2026-01-01T00:00:00+00:00",
    )


def _extract_pptx_structure(path) -> list[list[str]]:
    prs = Presentation(str(path))
    slides = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text for cell in row.cells))
        slides.append(texts)
    return slides


def _extract_docx_structure(path) -> dict:
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    return {"paragraphs": paragraphs, "tables": tables}


def _compare_to_golden(golden_path: Path, actual):
    if not golden_path.exists():
        golden_path.parent.mkdir(parents=True, exist_ok=True)
        golden_path.write_text(json.dumps(actual, indent=2), encoding="utf-8")
        pytest.skip(f"Golden file {golden_path} did not exist; created it from current output. Re-run to verify.")
    expected = json.loads(golden_path.read_text(encoding="utf-8"))
    assert actual == expected, (
        f"Rendered structure no longer matches {golden_path.name}. "
        "If this is an intentional template change, delete the golden file and re-run to re-baseline."
    )


def test_pptx_matches_golden_structure(tmp_path):
    report = _fixture_report()
    output_path = ValidationDeckBuilder().build(report, tmp_path / "report.pptx")
    actual = _extract_pptx_structure(output_path)
    _compare_to_golden(GOLDEN_DIR / "pptx_structure.json", actual)


def test_docx_matches_golden_structure(tmp_path):
    report = _fixture_report()
    output_path = ValidationWordReportBuilder().build(report, tmp_path / "report.docx")
    actual = _extract_docx_structure(output_path)
    _compare_to_golden(GOLDEN_DIR / "docx_structure.json", actual)
