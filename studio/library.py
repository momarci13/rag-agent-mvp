"""Shared knowledge library: regulations, concept papers and reference texts.

Attached documents are added to one persistent collection so every later
project can retrieve them. Each chunk keeps a locator (PDF page, DOCX or
Markdown heading) so agents can cite "document, p. 12" rather than a chunk id.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import threading
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Literal

from rag.ingest import chunk_text

from .frameworks import CORPUS_FOLDER_FRAMEWORK
from .schemas import Citation

DocKind = Literal["regulation", "concept_paper", "reference", "model_document"]
LIBRARY_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".tex"}


@dataclass
class LibraryDocument:
    doc_id: str
    title: str
    filename: str
    kind: str
    framework: str = ""
    sha256: str = ""
    segments: int = 0
    chunks: int = 0
    builtin: bool = False
    projects: list[str] = field(default_factory=list)


@dataclass
class Passage:
    chunk_id: str
    text: str
    citation: Citation
    score: float = 0.0
    doc_id: str = ""
    kind: str = ""

    def as_prompt(self) -> str:
        return f"[{self.chunk_id}] ({self.citation.text()})\n{self.text}"


# --------------------------------------------------------------------------
# Document segmentation with locators
# --------------------------------------------------------------------------

def _pdf_segments(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - damaged page
            text = ""
        if text.strip():
            out.append((f"p. {i}", text))
    return out


def _docx_segments(path: Path) -> list[tuple[str, str]]:
    """Walk paragraphs and tables in document order, splitting at headings."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out: list[tuple[str, str]] = []
    heading = "Introduction"
    buf: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            style = (para.style.name or "").lower() if para.style is not None else ""
            if not text:
                continue
            if style.startswith("heading") or style == "title":
                if buf:
                    out.append((f"section '{heading}'", "\n\n".join(buf)))
                    buf = []
                heading = text[:100]
            else:
                buf.append(text)
        elif tag == "tbl":
            table = Table(child, doc)
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
            if rows:
                buf.append("\n".join(rows))
    if buf:
        out.append((f"section '{heading}'", "\n\n".join(buf)))
    return out


_MD_HEADING = re.compile(r"^(#{1,4})\s+(.*)$", re.MULTILINE)


def _markdown_segments(text: str) -> list[tuple[str, str]]:
    matches = list(_MD_HEADING.finditer(text))
    if not matches:
        return [("", text)] if text.strip() else []
    out = []
    if matches[0].start() > 0 and text[: matches[0].start()].strip():
        out.append(("preamble", text[: matches[0].start()]))
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = text[m.end():end].strip()
        if body:
            out.append((f"section '{m.group(2).strip()[:100]}'", body))
    return out


def segment_document(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_segments(path)
    if suffix == ".docx":
        return _docx_segments(path)
    if suffix in (".md", ".txt", ".tex"):
        return _markdown_segments(path.read_text(encoding="utf-8", errors="ignore"))
    raise ValueError(f"Unsupported document type: {path.name}")


def document_text(path: Path, limit: int | None = None) -> str:
    """Whole text of a document with locator markers, for LLM review."""
    parts = [f"[{loc}]\n{txt}" if loc else txt for loc, txt in segment_document(path)]
    text = "\n\n".join(parts)
    return text[:limit] if limit else text


# --------------------------------------------------------------------------
# Library
# --------------------------------------------------------------------------

class Library:
    """Registry + hybrid retrieval over the shared document collection."""

    def __init__(self, root: Path, rag: Any, *, chunk_tokens: int = 256, overlap: int = 32) -> None:
        self.root = Path(root)
        self.files_dir = self.root / "files"
        self.files_dir.mkdir(parents=True, exist_ok=True)
        self.registry_path = self.root / "library.json"
        self.rag = rag
        self.chunk_tokens = chunk_tokens
        self.overlap = overlap
        self._lock = threading.RLock()
        self._docs: dict[str, LibraryDocument] = self._load_registry()

    # ---------- registry ----------
    def _load_registry(self) -> dict[str, LibraryDocument]:
        if not self.registry_path.exists():
            return {}
        raw = json.loads(self.registry_path.read_text(encoding="utf-8"))
        return {d["doc_id"]: LibraryDocument(**d) for d in raw.get("documents", [])}

    def _save_registry(self) -> None:
        payload = {"documents": [asdict(d) for d in self._docs.values()]}
        tmp = self.registry_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        tmp.replace(self.registry_path)

    def documents(self) -> list[LibraryDocument]:
        return sorted(self._docs.values(), key=lambda d: (d.builtin, d.title.lower()))

    def get(self, doc_id: str) -> LibraryDocument | None:
        return self._docs.get(doc_id)

    # ---------- ingestion ----------
    def add_document(
        self,
        path: Path,
        *,
        kind: str,
        framework: str = "",
        title: str | None = None,
        project_id: str | None = None,
        builtin: bool = False,
    ) -> LibraryDocument:
        path = Path(path)
        if path.suffix.lower() not in LIBRARY_EXTENSIONS:
            raise ValueError(f"{path.name}: supported types are {sorted(LIBRARY_EXTENSIONS)}")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        with self._lock:
            for doc in self._docs.values():
                if doc.sha256 == digest:
                    if project_id and project_id not in doc.projects:
                        doc.projects.append(project_id)
                        self._save_registry()
                    return doc

            doc_id = f"{'B' if builtin else 'D'}{digest[:10]}"
            stored = self.files_dir / f"{doc_id}{path.suffix.lower()}"
            if not builtin:
                shutil.copy2(path, stored)
            display = title or path.stem.replace("_", " ")
            segments = segment_document(path)
            chunks: list[dict[str, Any]] = []
            for s_idx, (locator, text) in enumerate(segments):
                for c_idx, chunk in enumerate(chunk_text(text, self.chunk_tokens, self.overlap)):
                    chunks.append({
                        "id": f"{doc_id}:{s_idx}:{c_idx}",
                        "text": chunk,
                        "meta": {
                            "source": doc_id,
                            "title": display[:200],
                            "locator": locator,
                            "kind": kind,
                            "framework": framework,
                        },
                    })
            existing = set(getattr(self.rag, "_ids", []))
            new_chunks = [c for c in chunks if c["id"] not in existing]
            for i in range(0, len(new_chunks), 64):
                self.rag.add(new_chunks[i : i + 64])
            doc = LibraryDocument(
                doc_id=doc_id,
                title=display,
                filename=path.name if builtin else stored.name,
                kind=kind,
                framework=framework,
                sha256=digest,
                segments=len(segments),
                chunks=len(chunks),
                builtin=builtin,
                projects=[project_id] if project_id else [],
            )
            self._docs[doc_id] = doc
            self._save_registry()
            return doc

    def ensure_builtin(self, corpus_dir: Path) -> int:
        """Ingest the bundled reference summaries once. Returns documents added."""
        added = 0
        if not corpus_dir.exists():
            return 0
        known = {d.sha256 for d in self._docs.values()}
        for path in sorted(corpus_dir.rglob("*.md")):
            if path.name.lower() == "readme.md":
                continue
            if hashlib.sha256(path.read_bytes()).hexdigest() in known:
                continue
            framework = CORPUS_FOLDER_FRAMEWORK.get(path.parent.name, "")
            title = f"{path.parent.name.replace('_', ' ').upper()}: {path.stem.replace('_', ' ')} (summary)"
            self.add_document(path, kind="reference", framework=framework, title=title, builtin=True)
            added += 1
        return added

    # ---------- retrieval ----------
    def _to_passage(self, hit: dict[str, Any]) -> Passage:
        meta = hit.get("meta") or {}
        doc = self._docs.get(meta.get("source", ""))
        title = meta.get("title") or (doc.title if doc else meta.get("source", "unknown"))
        return Passage(
            chunk_id=hit["id"],
            text=hit["text"],
            citation=Citation(source=title, locator=meta.get("locator", ""), chunk_id=hit["id"]),
            score=float(hit.get("score", 0.0)),
            doc_id=meta.get("source", ""),
            kind=meta.get("kind", ""),
        )

    def retrieve(
        self,
        query: str,
        *,
        k: int = 8,
        doc_ids: list[str] | None = None,
        frameworks: list[str] | None = None,
        token_budget: int = 6000,
    ) -> list[Passage]:
        if len(self.rag) == 0:
            return []
        hits: list[dict[str, Any]] = []
        if doc_ids:
            per_doc = max(2, k // max(1, len(doc_ids)) + 1)
            for doc_id in doc_ids:
                hits.extend(self.rag.retrieve(
                    query, k=per_doc, m=60, token_budget=token_budget,
                    metadata_filters={"source": doc_id},
                ))
        else:
            hits = self.rag.retrieve(query, k=k * 4, m=80, token_budget=token_budget * 3)
            if frameworks is not None:
                allowed = set(frameworks)
                hits = [
                    h for h in hits
                    if (h.get("meta") or {}).get("kind") != "reference"
                    or (h.get("meta") or {}).get("framework") in allowed
                ]
        seen: set[str] = set()
        passages: list[Passage] = []
        for h in sorted(hits, key=lambda x: x.get("score", 0.0), reverse=True):
            if h["id"] in seen:
                continue
            seen.add(h["id"])
            passages.append(self._to_passage(h))
            if len(passages) >= k:
                break
        return passages

    def chunk_citation(self, chunk_id: str) -> Citation | None:
        ids = getattr(self.rag, "_ids", [])
        try:
            idx = ids.index(chunk_id)
        except ValueError:
            return None
        meta = self.rag._metas[idx] if idx < len(self.rag._metas) else {}
        return Citation(source=meta.get("title", meta.get("source", "")), locator=meta.get("locator", ""), chunk_id=chunk_id)


def make_rag(cfg: dict[str, Any], collection: str, embedding_client: Any | None = None) -> Any:
    from rag.hybrid import LiteHybridRAG

    rag_cfg = cfg.get("rag", {})
    return LiteHybridRAG(
        db_path=_resolve(rag_cfg.get("db_path", "kb/chroma")),
        collection=collection,
        embedding_model=rag_cfg.get("embedding_model", "text-embedding-3-large"),
        embedding_client=embedding_client,
        alpha_dense=rag_cfg.get("alpha_dense", 0.6),
        query_expansion_enabled=False,
    )


def _resolve(path: str) -> str:
    from .config import ROOT

    p = Path(path)
    return str(p if p.is_absolute() else (ROOT / p).resolve())
