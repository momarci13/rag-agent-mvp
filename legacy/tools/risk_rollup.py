"""Executive rollup deck: aggregates multiple risk-validation reports into
one summary PPTX for a risk committee, e.g. "all Tier 1 credit models
validated this quarter". Reuses tools/pptx_report.py's slide shell
(blank_slide/add_banner) rather than duplicating it.
"""
from __future__ import annotations

import io
from collections import Counter
from pathlib import Path

from pptx.util import Inches, Pt

from agents.risk_schemas import ValidationReport
from tools.pptx_report import DISCLAIMER_BANNER, add_banner, blank_slide, new_presentation
from tools.report_charts import render_metric_chart

_RATING_TO_SEVERITY_BUCKET = {
    "compliant": "low", "low": "low", "not_applicable": "low",
    "partially_compliant": "medium", "medium": "medium",
    "non_compliant": "high", "high": "high",
    "unacceptable": "critical", "critical": "critical",
}


class ExecutiveRollupBuilder:
    """Builds a draft executive rollup deck from multiple ValidationReports."""

    def build(self, reports: list[ValidationReport], output_path: str | Path, *, title: str = "Risk Validation Rollup") -> Path:
        prs = new_presentation()

        self._title_slide(prs, title, reports)
        self._summary_table_slide(prs, reports)
        self._severity_distribution_slide(prs, reports)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))
        return output_path

    def _title_slide(self, prs, title: str, reports: list[ValidationReport]) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_banner(prs, slide)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), prs.slide_width - Inches(1.4), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.text = title
        tf.paragraphs[0].font.size = Pt(32)
        tf.paragraphs[0].font.bold = True

        meta_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.6), prs.slide_width - Inches(1.4), Inches(1.5))
        meta_tf = meta_box.text_frame
        meta_tf.text = f"{len(reports)} validation report(s) included"
        p = meta_tf.add_paragraph()
        p.text = DISCLAIMER_BANNER
        for para in meta_tf.paragraphs:
            para.font.size = Pt(16)

    def _summary_table_slide(self, prs, reports: list[ValidationReport]) -> None:
        slide = blank_slide(prs, "Included Reports")
        rows = len(reports) + 1
        table_shape = slide.shapes.add_table(rows, 4, Inches(0.4), Inches(1.3), Inches(12.5), Inches(0.4) * rows)
        table = table_shape.table
        widths = [Inches(5.5), Inches(2.5), Inches(2.5), Inches(2.0)]
        for i, w in enumerate(widths):
            table.columns[i].width = w
        for i, header in enumerate(["Entity", "Domain", "Rating", "Period"]):
            table.cell(0, i).text = header
        for row, report in enumerate(reports, start=1):
            table.cell(row, 0).text = report.entity_under_review
            table.cell(row, 1).text = report.domain.replace("_", " ").title()
            table.cell(row, 2).text = report.overall_rating
            table.cell(row, 3).text = report.reporting_period

    def _severity_distribution_slide(self, prs, reports: list[ValidationReport]) -> None:
        slide = blank_slide(prs, "Severity Distribution")
        counts = Counter(
            _RATING_TO_SEVERITY_BUCKET.get(r.overall_rating, "medium") for r in reports
        )
        data = {bucket: float(counts.get(bucket, 0)) for bucket in ("low", "medium", "high", "critical")}
        chart_bytes = render_metric_chart(data, "Reports by overall rating bucket")
        slide.shapes.add_picture(io.BytesIO(chart_bytes), Inches(3.0), Inches(1.5), width=Inches(7.0))
