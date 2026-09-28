"""DOCX (and derived PDF) builders for the modelling document and the
validation report.

Every table of results (tests, acceptance criteria, metrics, replication,
findings, requirement coverage) is filled from executed results, never from
LLM prose. The LLM only writes the narrative paragraphs.
"""
from __future__ import annotations

import datetime as dt
import io
from collections import Counter
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .schemas import (
    SEVERITY_ORDER,
    Finding,
    ModelerRound,
    ModelSpec,
    Project,
    Requirement,
    Round,
    TestCase,
    TestRun,
    ValidatorRound,
)

INK = RGBColor(0x1B, 0x1A, 0x17)
MUTED = RGBColor(0x66, 0x61, 0x55)
ALERT = RGBColor(0xA3, 0x20, 0x1C)
ACCENT = RGBColor(0x0A, 0x65, 0x59)

OUTCOME_LABEL = {
    "approved": "Approved",
    "approved_with_conditions": "Approved with conditions",
    "remediation_required": "Remediation required",
    "rejected": "Rejected",
    None: "Pending",
}


# --------------------------------------------------------------------------
# Low-level helpers
# --------------------------------------------------------------------------

def _new_document() -> Document:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for level, size in ((1, 15), (2, 12.5), (3, 11)):
        st = styles[f"Heading {level}"]
        st.font.size = Pt(size)
        st.font.color.rgb = INK
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.9)
    return doc


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc: Document, header: list[str], rows: Iterable[list[str]], widths: list[float] | None = None) -> None:
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, "EEECE6")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(text))
            run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def _body(doc: Document, text: str) -> None:
    """Paragraphs separated by blank lines; lines starting with '- ' become bullets."""
    for block in (text or "").split("\n\n"):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        if all(l.strip().startswith(("- ", "* ")) for l in lines):
            for l in lines:
                doc.add_paragraph(l.strip()[2:], style="List Bullet")
        else:
            doc.add_paragraph(" ".join(l.strip() for l in lines))


def _bullets(doc: Document, items: list[str], empty: str = "None stated.") -> None:
    if not items:
        p = doc.add_paragraph(empty)
        p.runs[0].italic = True
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _banner(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.color.rgb = ALERT


def _meta_line(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.color.rgb = MUTED
    run.font.size = Pt(9)


def render_equation(latex: str) -> bytes | None:
    """Render LaTeX math to PNG with matplotlib mathtext; None if unsupported."""
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex.strip().strip('$')}$", fontsize=13)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", pad_inches=0.05, transparent=False)
        plt.close(fig)
        return buf.getvalue()
    except Exception:  # noqa: BLE001 - unsupported mathtext constructs
        try:
            plt.close("all")
        except Exception:  # noqa: BLE001
            pass
        return None


def _equation(doc: Document, label: str, latex: str, explanation: str) -> None:
    png = render_equation(latex)
    if png:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(io.BytesIO(png), height=Inches(min(0.9, 0.35 + 0.1 * latex.count("\\frac"))))
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(latex)
        r.font.name = "Consolas"
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"({label})")
    r.italic = True
    r.font.size = Pt(9)
    if explanation:
        doc.add_paragraph(explanation)


def _code(doc: Document, text: str, max_lines: int = 400) -> None:
    lines = text.splitlines()
    if len(lines) > max_lines:
        lines = lines[:max_lines] + [f"... [{len(text.splitlines()) - max_lines} more lines in the package]"]
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(7.5)


def _test_status(test_id: str, run: TestRun | None) -> tuple[str, str]:
    if run is None:
        return "Not run", ""
    matches = [r for r in run.results if r.test_id == test_id]
    if not matches:
        return "Not implemented", ""
    order = {"error": 3, "failed": 2, "skipped": 1, "passed": 0}
    worst = max(matches, key=lambda r: order[r.outcome])
    label = {"passed": "Pass", "failed": "Fail", "error": "Error", "skipped": "Skipped"}[worst.outcome]
    if len(matches) > 1:
        label += f" ({sum(r.outcome == 'passed' for r in matches)}/{len(matches)} passed)"
    return label, worst.message.splitlines()[0][:160] if worst.message else ""


def _test_plan_table(doc: Document, plan: list[TestCase], run: TestRun | None) -> None:
    rows = []
    for tc in plan:
        status, note = _test_status(tc.test_id, run)
        rows.append([tc.test_id, tc.category, tc.name, tc.acceptance, status + (f": {note}" if note else "")])
    _table(doc, ["ID", "Category", "Test", "Acceptance rule", "Result"], rows, [0.6, 0.9, 1.7, 1.9, 1.6])


def _figures(doc: Document, figures: list[str], limit: int = 8) -> None:
    for path in figures[:limit]:
        try:
            doc.add_picture(path, width=Inches(5.6))
            cap = doc.add_paragraph(Path(path).stem.replace("_", " "))
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
        except Exception:  # noqa: BLE001 - skip unreadable images
            continue


def _requirement_lookup(reqs: list[Requirement]) -> dict[str, Requirement]:
    return {r.req_id: r for r in reqs}


def to_pdf(docx_path: Path, engine: str = "libreoffice") -> Path | None:
    from tools.docx_to_pdf import convert_docx_to_pdf

    try:
        return convert_docx_to_pdf(docx_path, docx_path.parent, engine=engine)
    except ValueError:
        return None


# --------------------------------------------------------------------------
# Modelling document
# --------------------------------------------------------------------------

NARRATIVE_SLOTS = {
    "methodology": ("methodology", "rationale", "approach"),
    "results": ("result", "performance", "interpretation"),
    "limitations": ("limitation", "model risk", "weakness"),
    "use": ("use", "implementation", "deployment", "guidance"),
}


def _slot_for(heading: str) -> str:
    h = heading.lower()
    for slot, keys in NARRATIVE_SLOTS.items():
        if any(k in h for k in keys):
            return slot
    return "other"


def build_modelling_document(
    project: Project,
    rnd: Round,
    package_dir: Path,
    output_path: Path,
    requirements: list[Requirement],
) -> Path:
    m: ModelerRound = rnd.modeler
    spec: ModelSpec = m.spec  # type: ignore[assignment]
    narrative = m.narrative
    doc = _new_document()

    _banner(doc, "DRAFT MODEL DOCUMENTATION: FOR INDEPENDENT VALIDATION")
    doc.add_heading(spec.title, level=0)
    _meta_line(doc, (
        f"Project {project.project_id} · Version {rnd.number} · "
        f"{dt.date.today().isoformat()} · Prepared by the Modeler agent (AI) for review by a qualified model owner"
    ))

    doc.add_heading("Document control", level=1)
    _table(doc, ["Field", "Value"], [
        ["Model category", f"{spec.category}{' / ' + spec.subcategory if spec.subcategory else ''}"],
        ["Version", f"{rnd.number} (development round {rnd.number})"],
        ["Code package", "package/ (model/api.py, pipeline.py, tests/)"],
        ["Test suite", _counts_text(m.tests)],
        ["Status", "Submitted for validation"],
    ], [1.8, 4.9])

    slots: dict[str, list] = {}
    if narrative:
        for sec in narrative.sections:
            slots.setdefault(_slot_for(sec.heading), []).append(sec)

    doc.add_heading("1. Executive summary", level=1)
    _body(doc, narrative.executive_summary if narrative else spec.objective)

    doc.add_heading("2. Purpose, scope and intended use", level=1)
    _table(doc, ["Item", "Description"], [
        ["Objective", spec.objective],
        ["Intended use", spec.intended_use],
        ["Target / output", spec.target],
        ["Unit of analysis", spec.unit_of_analysis or "Not specified"],
    ], [1.6, 5.1])

    doc.add_heading("3. Data", level=1)
    _body(doc, spec.data_description)
    tables = (project.data_profile or {}).get("tables", {})
    if tables:
        rows = []
        for name, t in tables.items():
            worst_missing = max((c.get("missing_pct", 0) for c in t["columns"]), default=0)
            rows.append([name, t["file"], f"{t['rows']:,}", str(len(t["columns"])), f"{worst_missing:.2f}%", str(t.get("duplicate_rows", 0))])
        _table(doc, ["Table", "File", "Rows", "Columns", "Max missing", "Duplicates"], rows)
    doc.add_heading("Data treatment", level=2)
    _bullets(doc, spec.data_treatment)
    if spec.features:
        doc.add_heading("Explanatory variables", level=2)
        _bullets(doc, spec.features)

    doc.add_heading("4. Methodology", level=1)
    _body(doc, spec.methodology_summary)
    if spec.equations:
        doc.add_heading("Model equations", level=2)
        for i, eq in enumerate(spec.equations, start=1):
            _equation(doc, eq.label or f"Eq. {i}", eq.latex, eq.explanation)
    doc.add_heading("Assumptions", level=2)
    _bullets(doc, spec.assumptions)
    doc.add_heading("Alternatives considered", level=2)
    _bullets(doc, spec.alternatives_considered)
    for sec in slots.get("methodology", []):
        doc.add_heading(sec.heading, level=2)
        _body(doc, sec.body)

    doc.add_heading("5. Implementation", level=1)
    doc.add_paragraph(
        "The model is delivered as a Python package. pipeline.py runs the full estimation and "
        "evaluation from the input data and writes results/metrics.json; model/api.py exposes "
        "load_data, fit, predict, evaluate and run; tests/ holds the automated test suite that "
        "implements the test plan in section 6."
    )
    _table(doc, ["File", "Lines"], [[f, str(_line_count(package_dir / f))] for f in m.files], [5.2, 1.5])
    doc.add_paragraph("Reproduce with:  python pipeline.py --data-dir <data> --out-dir results   and   python -m pytest tests")

    doc.add_heading("6. Testing framework", level=1)
    doc.add_paragraph(
        "Each planned test is implemented as an automated pytest test whose name carries the test ID. "
        "Results below are taken from the executed test run, not from the development narrative."
    )
    _test_plan_table(doc, spec.test_plan, m.tests)
    extra = [r for r in (m.tests.results if m.tests else []) if r.test_id is None]
    if extra:
        doc.add_paragraph(f"{len(extra)} additional unit tests without a plan ID also ran: "
                          f"{sum(r.outcome == 'passed' for r in extra)} passed.")

    doc.add_heading("7. Results", level=1)
    metrics = m.pipeline.metrics if m.pipeline else {}
    if spec.acceptance_criteria:
        doc.add_heading("Acceptance criteria", level=2)
        rows = []
        for ac in spec.acceptance_criteria:
            val = metrics.get(ac.metric)
            status = "Missing" if val is None else ("Pass" if ac.passes(val) else "Fail")
            rows.append([ac.metric + (" (primary)" if ac.primary else ""), ac.describe(),
                         "n/a" if val is None else f"{val:.6g}", status])
        _table(doc, ["Metric", "Criterion", "Value", "Result"], rows, [1.9, 2.3, 1.2, 1.3])
    if metrics:
        doc.add_heading("All reported metrics", level=2)
        _table(doc, ["Metric", "Value"], [[k, f"{v:.6g}"] for k, v in sorted(metrics.items())], [4.2, 2.5])
    if m.pipeline and m.pipeline.error:
        _banner(doc, "The pipeline reported an error: " + m.pipeline.error[:500])
    _figures(doc, m.pipeline.figures if m.pipeline else [])
    for sec in slots.get("results", []):
        doc.add_heading(sec.heading, level=2)
        _body(doc, sec.body)

    doc.add_heading("8. Limitations and model risk", level=1)
    _bullets(doc, spec.limitations)
    for sec in slots.get("limitations", []):
        doc.add_heading(sec.heading, level=2)
        _body(doc, sec.body)

    doc.add_heading("9. Requirements addressed", level=1)
    lookup = _requirement_lookup(requirements)
    rows = []
    for rid in spec.requirement_ids:
        req = lookup.get(rid)
        if req:
            rows.append([rid, req.text, req.citation.text() if req.citation else req.framework])
        else:
            rows.append([rid, "(not in the current checklist)", ""])
    if rows:
        _table(doc, ["ID", "Requirement", "Source"], rows, [0.9, 4.0, 1.8])
    else:
        doc.add_paragraph("No requirement mapping provided.")

    doc.add_heading("10. Ongoing monitoring", level=1)
    _bullets(doc, spec.monitoring_plan)
    for sec in slots.get("use", []) + slots.get("other", []):
        doc.add_heading(sec.heading, level=2)
        _body(doc, sec.body)

    if m.responses:
        doc.add_heading("11. Response to validation findings", level=1)
        _table(doc, ["Finding", "Action", "Explanation", "Changes"], [
            [r.finding_id, r.action.replace("_", " "), r.explanation, "; ".join(r.changes)] for r in m.responses
        ], [1.2, 1.0, 2.8, 1.7])

    doc.add_page_break()
    doc.add_heading("Appendix A. Test execution log", level=1)
    if m.tests:
        doc.add_paragraph(_counts_text(m.tests))
        for r in m.tests.results:
            if r.outcome in ("failed", "error"):
                doc.add_paragraph(f"{r.nodeid}: {r.outcome}", style="List Bullet")
                if r.message:
                    _code(doc, r.message, max_lines=15)
        if m.tests.collection_error:
            _code(doc, m.tests.collection_error, max_lines=40)
    doc.add_heading("Appendix B. Source code", level=1)
    for rel in m.files:
        if rel.endswith(".py"):
            doc.add_heading(rel, level=3)
            _code(doc, (package_dir / rel).read_text(encoding="utf-8", errors="replace"))

    _meta_line(doc, project.disclaimer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _counts_text(run: TestRun | None) -> str:
    if run is None:
        return "Not run"
    c = run.counts()
    text = f"{c['passed']} passed, {c['failed']} failed, {c['error']} errors, {c['skipped']} skipped"
    if run.collection_error:
        text += " (test collection error)"
    return text


def _line_count(path: Path) -> int:
    try:
        return len(path.read_text(encoding="utf-8", errors="replace").splitlines())
    except OSError:
        return 0


# --------------------------------------------------------------------------
# Validation report
# --------------------------------------------------------------------------

def build_validation_report(
    project: Project,
    rnd: Round,
    output_path: Path,
    *,
    subject_title: str,
    attached_titles: list[str],
) -> Path:
    v: ValidatorRound = rnd.validator
    doc = _new_document()
    _banner(doc, "DRAFT VALIDATION REPORT: REQUIRES SIGN-OFF BY A QUALIFIED VALIDATOR")
    doc.add_heading(f"Validation report: {subject_title}", level=0)
    _meta_line(doc, (
        f"Project {project.project_id} · Validation round {rnd.number} · {dt.date.today().isoformat()} · "
        "Prepared by the Validator agent (AI)"
    ))

    sev = Counter(f.severity for f in v.findings if f.status == "open")
    _table(doc, ["Item", "Value"], [
        ["Validation outcome (rule-based)", OUTCOME_LABEL[v.outcome]],
        ["Open findings", ", ".join(f"{sev[s]} {s}" for s in ("critical", "high", "medium", "low", "observation") if sev[s]) or "None"],
        ["Closed in this round", str(sum(1 for f in v.findings if f.status != "open" and f.closed_round == rnd.number))],
        ["Mode", "Develop and validate" if project.input.mode == "develop_and_validate" else "Validation of an external model"],
        ["Frameworks", ", ".join(project.input.frameworks) or "Generic practice only"],
    ], [2.4, 4.3])

    doc.add_heading("1. Conclusion", level=1)
    _body(doc, v.conclusion or "Conclusion pending.")
    doc.add_paragraph(
        "The outcome is set by rules, not by the AI: any open critical finding means Rejected; any open high "
        "finding means Remediation required; open medium findings mean Approved with conditions; otherwise Approved."
    )

    doc.add_heading("2. Scope and approach", level=1)
    steps = [
        "Replication: the delivered pipeline was re-run from a clean copy and every reported metric compared.",
        "Re-execution of the developer's test suite and a check that every planned test exists and passes.",
        "Document and conceptual soundness review against the requirement checklist below.",
        f"Independent tests designed and run by the Validator ({len(v.independent_plan)} planned).",
    ]
    if v.challenger is not None:
        steps.append("Independent challenger model built from the brief and attached concept papers, then compared with the champion.")
    _bullets(doc, steps)
    if attached_titles:
        doc.add_heading("Documents relied on", level=2)
        _bullets(doc, attached_titles)

    doc.add_heading("3. Findings", level=1)
    ordered = sorted(v.findings, key=lambda f: (f.status != "open", -SEVERITY_ORDER[f.severity], f.finding_id))
    if ordered:
        _table(doc, ["ID", "Severity", "Area", "Title", "Status"], [
            [f.finding_id, f.severity, f.area, f.title, f.status.replace("_", " ")] for f in ordered
        ], [1.2, 0.8, 1.3, 2.6, 0.8])
        open_f = [f for f in ordered if f.status == "open"]
        for f in open_f:
            _finding_detail(doc, f)
        closed = [f for f in ordered if f.status != "open"]
        if closed:
            doc.add_heading("Closed and resolved findings", level=2)
            _table(doc, ["ID", "Title", "Closed", "Developer response"], [
                [f.finding_id, f.title, f"round {f.closed_round}" if f.closed_round else f.status,
                 (f.modeler_response.action.replace("_", " ") + ": " + f.modeler_response.explanation) if f.modeler_response else ""]
                for f in closed
            ], [1.2, 2.2, 0.8, 2.5])
    else:
        doc.add_paragraph("No findings.")

    doc.add_heading("4. Requirement coverage", level=1)
    assess = {a.req_id: a for a in (v.review.requirement_assessments if v.review else [])}
    rows = []
    for req in v.requirements:
        a = assess.get(req.req_id)
        rows.append([
            req.req_id, req.text, req.citation.text() if req.citation else req.framework,
            (a.status if a else "not_assessed").replace("_", " "), a.evidence[:200] if a else "",
        ])
    _table(doc, ["ID", "Requirement", "Source", "Status", "Evidence"], rows, [0.8, 2.3, 1.3, 0.8, 1.5])

    doc.add_heading("5. Replication", level=1)
    rep = v.replication
    if rep:
        doc.add_paragraph(("Reproduced. " if rep.reproduced else "Not reproduced. ") + rep.note)
        if rep.compared:
            _table(doc, ["Metric", "Reported", "Replicated", "Abs. difference"], [
                [k, f"{d.get('reported', float('nan')):.6g}", f"{d.get('replicated', float('nan')):.6g}", f"{d.get('abs_diff', float('nan')):.3g}"]
                for k, d in sorted(rep.compared.items())
            ])
    else:
        doc.add_paragraph("Not performed.")

    doc.add_heading("6. Developer test suite (re-run by the Validator)", level=1)
    doc.add_paragraph(_counts_text(v.modeler_tests_rerun))
    spec = rnd.modeler.spec
    if spec and spec.test_plan:
        _test_plan_table(doc, spec.test_plan, v.modeler_tests_rerun)

    doc.add_heading("7. Independent tests", level=1)
    doc.add_paragraph(_counts_text(v.independent_tests))
    if v.independent_plan:
        _test_plan_table(doc, v.independent_plan, v.independent_tests)
        for tc in v.independent_plan:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{tc.test_id} {tc.name}: ").bold = True
            p.add_run(f"{tc.objective} Method: {tc.method}")

    doc.add_heading("8. Challenger model", level=1)
    ch = v.challenger
    if ch is None:
        doc.add_paragraph("No challenger was built in this round.")
    else:
        _body(doc, ch.description or "")
        if ch.error:
            _banner(doc, "Challenger error: " + ch.error[:400])
        keys = sorted(set(ch.champion_metrics) & set(ch.challenger_metrics))
        if keys:
            _table(doc, ["Metric", "Champion", "Challenger"], [
                [k + (" (primary)" if k == ch.primary_metric else ""), f"{ch.champion_metrics[k]:.6g}", f"{ch.challenger_metrics[k]:.6g}"]
                for k in keys
            ])
        if ch.challenger_better is not None:
            doc.add_paragraph("The challenger materially outperforms the champion on the primary metric."
                              if ch.challenger_better else "The champion performs at least as well as the challenger on the primary metric.")

    doc.add_heading("9. Finding history", level=1)
    rows = []
    for past in project.rounds:
        for f in past.validator.findings:
            rows.append([str(past.number), f.finding_id, f.severity, f.status.replace("_", " "),
                         f.modeler_response.action.replace("_", " ") if f.modeler_response else ""])
    if rows:
        _table(doc, ["Round", "Finding", "Severity", "Status", "Developer response"], rows)

    doc.add_heading("10. Sign-off", level=1)
    sign_rows = [[s.name, s.role, s.decision, s.at, s.comment] for s in project.signoffs] or [["", "", "", "", ""]]
    _table(doc, ["Name", "Role", "Decision", "Date", "Comment"], sign_rows)

    _meta_line(doc, project.disclaimer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _finding_detail(doc: Document, f: Finding) -> None:
    doc.add_heading(f"{f.finding_id}: {f.title}", level=3)
    _meta_line(doc, f"Severity {f.severity} · {f.area} · Source {f.source.replace('_', ' ')} · Raised in round {f.raised_round}"
               + (f" · Closed in round {f.closed_round}" if f.closed_round else ""))
    _body(doc, f.description)
    if f.evidence:
        doc.add_paragraph("Evidence:")
        _bullets(doc, f.evidence[:8])
    if f.citations:
        doc.add_paragraph("References: " + "; ".join(c.text() for c in f.citations))
    if f.recommendation:
        p = doc.add_paragraph()
        p.add_run("Recommendation: ").bold = True
        p.add_run(f.recommendation)
    if f.modeler_response:
        p = doc.add_paragraph()
        p.add_run(f"Developer response ({f.modeler_response.action.replace('_', ' ')}): ").bold = True
        p.add_run(f.modeler_response.explanation)
