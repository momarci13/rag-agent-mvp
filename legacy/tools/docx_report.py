"""DOCX validation report builder.

Builds a draft Word document from a single, domain-agnostic
:class:`agents.risk_schemas.ValidationReport`. This is the source document
for the derived PDF -- see tools/docx_to_pdf.py -- so table/section content
must not drift from tools/pptx_report.py's deck.
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from agents.risk_schemas import ValidationReport
from tools.pptx_report import _shared_numeric_trend_series
from tools.report_charts import render_metric_chart

DISCLAIMER_BANNER = "DRAFT -- NOT A REGULATORY SUBMISSION"

_MAX_TREND_CHARTS = 3


class ValidationWordReportBuilder:
    """Builds a draft Word validation report from a ValidationReport."""

    def build(
        self,
        report: ValidationReport,
        output_path: str | Path,
        *,
        history: list[ValidationReport] | None = None,
    ) -> Path:
        doc = Document()

        banner = doc.add_paragraph()
        banner_run = banner.add_run(DISCLAIMER_BANNER)
        banner_run.bold = True
        banner_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

        doc.add_heading(report.title, level=0)
        meta = doc.add_paragraph()
        meta.add_run(
            f"Entity under review: {report.entity_under_review}\n"
            f"Reporting period: {report.reporting_period}\n"
            f"Domain: {report.domain.replace('_', ' ').title()}\n"
            f"Generated: {report.generated_at}\n"
            f"Prepared by: {report.prepared_by}"
        )

        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(report.overall_conclusion or "(pending)")

        if report.recommendation and report.recommendation != "not_a_recommendation":
            doc.add_heading("Recommendation & Conditions", level=1)
            rec_p = doc.add_paragraph()
            rec_run = rec_p.add_run(f"Recommendation: {report.recommendation.replace('_', ' ').upper()}")
            rec_run.bold = True
            for cond in report.conditions:
                doc.add_paragraph(cond, style="List Bullet")
            doc.add_paragraph(
                "This recommendation is derived deterministically from the overall rating and the "
                "deterministic gate result. It is a drafting aid, not a supervisory decision."
            )

        if report.follow_up_on_prior_findings:
            doc.add_heading("Follow-up on Prior Findings", level=1)
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, header in enumerate(["Reference", "Area", "Status", "Note"]):
                hdr[i].text = header
            for pf in report.follow_up_on_prior_findings:
                row = table.add_row().cells
                row[0].text = pf.finding_reference
                row[1].text = pf.area
                row[2].text = pf.status
                row[3].text = pf.note

        if report.validation_sample or report.materiality_rationale or report.deviations_from_policy:
            doc.add_heading("Validation Sample & Materiality", level=1)
            if report.validation_sample:
                doc.add_heading("Validation sample", level=2)
                doc.add_paragraph(report.validation_sample)
            if report.materiality_rationale:
                doc.add_heading("Materiality rationale", level=2)
                doc.add_paragraph(report.materiality_rationale)
            if report.deviations_from_policy:
                doc.add_heading("Deviations from policy", level=2)
                for dev in report.deviations_from_policy:
                    doc.add_paragraph(dev, style="List Bullet")

        doc.add_heading("Scope & Methodology", level=1)
        doc.add_heading("Scope", level=2)
        doc.add_paragraph(report.scope)
        doc.add_heading("Methodology", level=2)
        doc.add_paragraph(report.methodology)

        doc.add_heading("Model / Exposure Overview", level=1)
        overview = doc.add_paragraph()
        overview.add_run(
            f"Entity under review: {report.entity_under_review}\n"
            f"Reporting period: {report.reporting_period}\n"
        )
        for key, value in report.quantitative_results.items():
            doc.add_paragraph(f"{key}: {value}", style="List Bullet")

        doc.add_heading("Validation Activities Checklist", level=1)
        areas: dict[str, str] = {}
        for f in report.findings:
            if f.area not in areas or f.verdict == "non_compliant":
                areas[f.area] = f.verdict
        if areas:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text = "Area", "Verdict"
            for area, verdict in areas.items():
                row = table.add_row().cells
                row[0].text, row[1].text = area, verdict
        else:
            doc.add_paragraph("No validation areas assessed.")

        doc.add_heading("Findings & Severity Ratings", level=1)
        if report.findings:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, header in enumerate(["Area", "Verdict", "Severity", "Description"]):
                hdr[i].text = header
            for f in report.findings:
                row = table.add_row().cells
                row[0].text = f.area
                row[1].text = f.verdict
                row[2].text = f.severity
                row[3].text = f.description
        else:
            doc.add_paragraph("No findings recorded.")

        doc.add_heading("Quantitative Test Results", level=1)
        if report.quantitative_results:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text = "Metric", "Value"
            for key, value in report.quantitative_results.items():
                row = table.add_row().cells
                row[0].text, row[1].text = str(key), str(value)
            numeric_results = {
                k: v for k, v in report.quantitative_results.items()
                if isinstance(v, (int, float)) and not isinstance(v, bool)
            }
            if numeric_results:
                chart_bytes = render_metric_chart(numeric_results, "Current Metrics")
                doc.add_picture(io.BytesIO(chart_bytes), width=Inches(6.0))
        else:
            doc.add_paragraph("No quantitative results were supplied.")

        if history:
            series_by_key = _shared_numeric_trend_series(report.quantitative_results, history)
            if series_by_key:
                doc.add_heading("Metric Trends", level=1)
                for key, series in list(series_by_key.items())[:_MAX_TREND_CHARTS]:
                    doc.add_heading(key, level=2)
                    labeled = {f"cycle {i + 1}": value for i, value in enumerate(series)}
                    chart_bytes = render_metric_chart(labeled, f"{key} across validation cycles")
                    doc.add_picture(io.BytesIO(chart_bytes), width=Inches(6.0))

        doc.add_heading("Reviewer Notes", level=1)
        doc.add_paragraph(
            "Space for the human validator to add remarks per finding before sign-off. "
            "This is a plain editable Word table cell, not a native Word comment."
        )
        if report.findings:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text = "Finding", "Reviewer note"
            for f in report.findings:
                row = table.add_row().cells
                row[0].text = f"[{f.severity}] {f.area}: {f.description[:120]}"
                row[1].text = ""  # intentionally empty -- editable by the validator
        else:
            doc.add_paragraph("No findings to annotate.")

        doc.add_heading("Recommendations & Remediation Plan", level=1)
        actionable = [f for f in report.findings if f.recommendation]
        if actionable:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, header in enumerate(["Recommendation", "Owner", "Deadline (days)"]):
                hdr[i].text = header
            for f in actionable:
                row = table.add_row().cells
                row[0].text = f.recommendation
                row[1].text = f.owner or "Unassigned"
                row[2].text = str(f.remediation_deadline_days) if f.remediation_deadline_days is not None else "n/a"
        else:
            doc.add_paragraph("No outstanding recommendations.")

        doc.add_heading("Overall Validation Conclusion", level=1)
        rating_p = doc.add_paragraph()
        rating_run = rating_p.add_run(f"Overall rating: {report.overall_rating.upper()}")
        rating_run.bold = True
        rating_run.font.size = Pt(14)
        doc.add_paragraph(report.overall_conclusion or "(pending)")

        doc.add_heading("Sign-off", level=1)
        doc.add_paragraph(f"Prepared by: {report.prepared_by}")
        if report.preparer:
            doc.add_paragraph(f"Case file prepared by: {report.preparer} (may not sign off this report)")
        doc.add_paragraph(f"Required sign-offs: {report.required_signoffs}")
        if report.signoffs:
            for signoff in report.signoffs:
                role_suffix = f" ({signoff.role})" if signoff.role else ""
                doc.add_paragraph(f"Signed off by {signoff.by}{role_suffix} at {signoff.at}", style="List Bullet")
        else:
            doc.add_paragraph("PENDING -- human validator sign-off required.")
        disclaimer_p = doc.add_paragraph()
        disclaimer_run = disclaimer_p.add_run(report.disclaimer)
        disclaimer_run.italic = True

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path
