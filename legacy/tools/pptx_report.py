"""PPTX validation deck builder.

Builds a draft PowerPoint deck from a single, domain-agnostic
:class:`agents.risk_schemas.ValidationReport`. Every slide carries a visible
"DRAFT -- not a regulatory submission" banner, and the sign-off slide is
regenerated (not just re-labeled) once the report has been signed off --
see agents/risk_validation_team.py::RiskValidationOrchestrator.execute.

``blank_slide``/``add_banner`` are module-level (not just builder methods) so
tools/risk_rollup.py's executive rollup deck can reuse the same slide shell
without duplicating it.
"""
from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from agents.risk_schemas import ValidationFinding, ValidationReport
from tools.report_charts import render_metric_chart

DISCLAIMER_BANNER = "DRAFT -- NOT A REGULATORY SUBMISSION"

_SEVERITY_COLOR = {
    "critical": RGBColor(0xC0, 0x00, 0x00),
    "high": RGBColor(0xE0, 0x6C, 0x0C),
    "medium": RGBColor(0xBF, 0x8F, 0x00),
    "low": RGBColor(0x54, 0x82, 0x35),
    "observation": RGBColor(0x59, 0x59, 0x59),
}

_MAX_FINDING_ROWS_PER_SLIDE = 8
_MAX_TREND_CHARTS = 3


def add_banner(prs: Presentation, slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.3), prs.slide_height - Inches(0.4), prs.slide_width - Inches(0.6), Inches(0.3))
    tf = box.text_frame
    tf.text = DISCLAIMER_BANNER
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    run.font.bold = True


def blank_slide(prs: Presentation, heading: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    add_banner(prs, slide)
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), prs.slide_width - Inches(1.0), Inches(0.7))
    tf = title_box.text_frame
    tf.text = heading
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    return slide


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _shared_numeric_trend_series(
    current: dict, history: list[ValidationReport],
) -> dict[str, list[float]]:
    """Top-level numeric keys present in ``current`` and at least one prior
    cycle's quantitative_results, each mapped to its value series
    (oldest..newest, current last)."""
    series: dict[str, list[float]] = {}
    for key, value in current.items():
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            continue
        prior = [
            r.quantitative_results[key]
            for r in history
            if isinstance(r.quantitative_results.get(key), (int, float))
            and not isinstance(r.quantitative_results.get(key), bool)
        ]
        if prior:
            series[key] = [float(v) for v in prior] + [float(value)]
    return series


class ValidationDeckBuilder:
    """Builds a draft PPTX validation deck from a ValidationReport."""

    def build(
        self,
        report: ValidationReport,
        output_path: str | Path,
        *,
        history: list[ValidationReport] | None = None,
    ) -> Path:
        prs = new_presentation()

        self._title_slide(prs, report)
        self._text_slide(prs, "Executive Summary", report.overall_conclusion or "(pending)")
        self._recommendation_slide(prs, report)
        self._prior_findings_slide(prs, report)
        self._validation_sample_slide(prs, report)
        self._text_slide(
            prs, "Scope & Methodology",
            f"Scope:\n{report.scope}\n\nMethodology:\n{report.methodology}",
        )
        self._overview_slide(prs, report)
        self._activities_slide(prs, report)
        self._findings_slides(prs, report)
        self._quantitative_slide(prs, report)
        if history:
            self._trend_slides(prs, report, history)
        self._recommendations_slide(prs, report)
        self._conclusion_slide(prs, report)
        self._signoff_slide(prs, report)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))
        return output_path

    # ---------- slide helpers ----------

    def _blank_slide(self, prs: Presentation, heading: str):
        return blank_slide(prs, heading)

    def _banner(self, prs: Presentation, slide) -> None:
        add_banner(prs, slide)

    def _title_slide(self, prs: Presentation, report: ValidationReport) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_banner(prs, slide)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), prs.slide_width - Inches(1.4), Inches(1.5))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.text = report.title
        tf.paragraphs[0].font.size = Pt(32)
        tf.paragraphs[0].font.bold = True

        meta_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.8), prs.slide_width - Inches(1.4), Inches(2.0))
        meta_tf = meta_box.text_frame
        meta_tf.word_wrap = True
        lines = [
            f"Entity under review: {report.entity_under_review}",
            f"Reporting period: {report.reporting_period}",
            f"Domain: {report.domain.replace('_', ' ').title()}",
            f"Generated: {report.generated_at}",
            f"Prepared by: {report.prepared_by}",
        ]
        meta_tf.text = lines[0]
        for line in lines[1:]:
            p = meta_tf.add_paragraph()
            p.text = line
        for p in meta_tf.paragraphs:
            p.font.size = Pt(14)

    def _text_slide(self, prs: Presentation, heading: str, body: str) -> None:
        slide = self._blank_slide(prs, heading)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), prs.slide_width - Inches(1.0), Inches(5.5))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = body
        for p in tf.paragraphs:
            p.font.size = Pt(16)

    def _recommendation_slide(self, prs: Presentation, report: ValidationReport) -> None:
        """Rendered only when the derived recommendation is committal -- keeps
        the deck unchanged for reports that carry the default."""
        if not report.recommendation or report.recommendation == "not_a_recommendation":
            return
        lines = [
            f"Recommendation: {report.recommendation.replace('_', ' ').upper()}",
            "",
            (
                "Derived deterministically from the overall rating and the gate result; "
                "a drafting aid, not a supervisory decision."
            ),
        ]
        if report.conditions:
            lines.append("")
            lines.append("Conditions:")
            lines.extend(f"  - {c}" for c in report.conditions)
        self._text_slide(prs, "Recommendation & Conditions", "\n".join(lines))

    def _prior_findings_slide(self, prs: Presentation, report: ValidationReport) -> None:
        rows_data = report.follow_up_on_prior_findings
        if not rows_data:
            return
        slide = self._blank_slide(prs, "Follow-up on Prior Findings")
        rows = len(rows_data) + 1
        table_shape = slide.shapes.add_table(rows, 4, Inches(0.4), Inches(1.3), Inches(12.5), Inches(0.5) * rows)
        table = table_shape.table
        for i, w in enumerate((Inches(3.4), Inches(2.6), Inches(1.6), Inches(4.9))):
            table.columns[i].width = w
        for i, header in enumerate(["Reference", "Area", "Status", "Note"]):
            table.cell(0, i).text = header
        for row, pf in enumerate(rows_data, start=1):
            table.cell(row, 0).text = pf.finding_reference
            table.cell(row, 1).text = pf.area
            table.cell(row, 2).text = pf.status
            table.cell(row, 3).text = pf.note

    def _validation_sample_slide(self, prs: Presentation, report: ValidationReport) -> None:
        if not (report.validation_sample or report.materiality_rationale or report.deviations_from_policy):
            return
        parts: list[str] = []
        if report.validation_sample:
            parts.append(f"Validation sample:\n{report.validation_sample}")
        if report.materiality_rationale:
            parts.append(f"Materiality rationale:\n{report.materiality_rationale}")
        if report.deviations_from_policy:
            parts.append("Deviations from policy:\n" + "\n".join(f"  - {d}" for d in report.deviations_from_policy))
        self._text_slide(prs, "Validation Sample & Materiality", "\n\n".join(parts))

    def _overview_slide(self, prs: Presentation, report: ValidationReport) -> None:
        body = (
            f"Entity under review: {report.entity_under_review}\n"
            f"Reporting period: {report.reporting_period}\n"
            f"Domain: {report.domain.replace('_', ' ').title()}\n\n"
            f"Case file summary (as submitted):\n"
        )
        for key, value in report.quantitative_results.items():
            body += f"  - {key}: {value}\n"
        self._text_slide(prs, "Model / Exposure Overview", body)

    def _activities_slide(self, prs: Presentation, report: ValidationReport) -> None:
        slide = self._blank_slide(prs, "Validation Activities Checklist")
        areas: dict[str, str] = {}
        for f in report.findings:
            if f.area not in areas or f.verdict == "non_compliant":
                areas[f.area] = f.verdict
        rows = max(1, len(areas)) + 1
        table_shape = slide.shapes.add_table(rows, 2, Inches(0.5), Inches(1.3), Inches(12.0), Inches(0.5) * rows)
        table = table_shape.table
        table.columns[0].width = Inches(9.0)
        table.columns[1].width = Inches(3.0)
        table.cell(0, 0).text = "Area"
        table.cell(0, 1).text = "Verdict"
        if not areas:
            table.cell(1, 0).text = "No validation areas assessed"
            table.cell(1, 1).text = "n/a"
        else:
            for i, (area, verdict) in enumerate(areas.items(), start=1):
                table.cell(i, 0).text = area
                table.cell(i, 1).text = verdict

    def _findings_slides(self, prs: Presentation, report: ValidationReport) -> None:
        findings = report.findings
        if not findings:
            self._text_slide(prs, "Findings & Severity Ratings", "No findings recorded.")
            return
        for start in range(0, len(findings), _MAX_FINDING_ROWS_PER_SLIDE):
            chunk = findings[start:start + _MAX_FINDING_ROWS_PER_SLIDE]
            heading = "Findings & Severity Ratings"
            if len(findings) > _MAX_FINDING_ROWS_PER_SLIDE:
                heading += f" ({start // _MAX_FINDING_ROWS_PER_SLIDE + 1})"
            self._findings_table_slide(prs, heading, chunk)

    def _findings_table_slide(self, prs: Presentation, heading: str, findings: list[ValidationFinding]) -> None:
        slide = self._blank_slide(prs, heading)
        rows = len(findings) + 1
        table_shape = slide.shapes.add_table(rows, 4, Inches(0.4), Inches(1.3), Inches(12.5), Inches(0.5) * rows)
        table = table_shape.table
        widths = [Inches(2.6), Inches(1.6), Inches(1.2), Inches(7.1)]
        for i, w in enumerate(widths):
            table.columns[i].width = w
        for i, header in enumerate(["Area", "Verdict", "Severity", "Description"]):
            table.cell(0, i).text = header
        for row, f in enumerate(findings, start=1):
            table.cell(row, 0).text = f.area
            table.cell(row, 1).text = f.verdict
            cell = table.cell(row, 2)
            cell.text = f.severity
            color = _SEVERITY_COLOR.get(f.severity)
            if color is not None and cell.text_frame.paragraphs[0].runs:
                cell.text_frame.paragraphs[0].runs[0].font.color.rgb = color
            table.cell(row, 3).text = f.description[:280]

    def _quantitative_slide(self, prs: Presentation, report: ValidationReport) -> None:
        slide = self._blank_slide(prs, "Quantitative Test Results")
        results = report.quantitative_results
        if not results:
            box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(11.0), Inches(1.0))
            box.text_frame.text = "No quantitative results were supplied."
            return
        rows = len(results) + 1
        table_shape = slide.shapes.add_table(rows, 2, Inches(0.5), Inches(1.3), Inches(5.5), Inches(0.4) * rows)
        table = table_shape.table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        for i, (key, value) in enumerate(results.items(), start=1):
            table.cell(i, 0).text = str(key)
            table.cell(i, 1).text = str(value)

        numeric_results = {
            k: v for k, v in results.items() if isinstance(v, (int, float)) and not isinstance(v, bool)
        }
        if numeric_results:
            chart_bytes = render_metric_chart(numeric_results, "Current Metrics")
            slide.shapes.add_picture(io.BytesIO(chart_bytes), Inches(6.5), Inches(1.3), width=Inches(6.3))

    def _trend_slides(self, prs: Presentation, report: ValidationReport, history: list[ValidationReport]) -> None:
        series_by_key = _shared_numeric_trend_series(report.quantitative_results, history)
        for key, series in list(series_by_key.items())[:_MAX_TREND_CHARTS]:
            slide = self._blank_slide(prs, f"Metric Trend: {key}")
            labeled = {f"cycle {i + 1}": value for i, value in enumerate(series)}
            chart_bytes = render_metric_chart(labeled, f"{key} across validation cycles")
            slide.shapes.add_picture(io.BytesIO(chart_bytes), Inches(2.5), Inches(1.5), width=Inches(8.0))

    def _recommendations_slide(self, prs: Presentation, report: ValidationReport) -> None:
        actionable = [f for f in report.findings if f.recommendation]
        slide = self._blank_slide(prs, "Recommendations & Remediation Plan")
        if not actionable:
            box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(11.0), Inches(1.0))
            box.text_frame.text = "No outstanding recommendations."
            return
        rows = len(actionable) + 1
        table_shape = slide.shapes.add_table(rows, 3, Inches(0.4), Inches(1.3), Inches(12.5), Inches(0.5) * rows)
        table = table_shape.table
        widths = [Inches(7.0), Inches(2.75), Inches(2.75)]
        for i, w in enumerate(widths):
            table.columns[i].width = w
        for i, header in enumerate(["Recommendation", "Owner", "Deadline (days)"]):
            table.cell(0, i).text = header
        for row, f in enumerate(actionable, start=1):
            table.cell(row, 0).text = f.recommendation[:280]
            table.cell(row, 1).text = f.owner or "Unassigned"
            table.cell(row, 2).text = str(f.remediation_deadline_days) if f.remediation_deadline_days is not None else "n/a"

    def _conclusion_slide(self, prs: Presentation, report: ValidationReport) -> None:
        slide = self._blank_slide(prs, "Overall Validation Conclusion")
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), prs.slide_width - Inches(1.4), Inches(1.0))
        tf = box.text_frame
        tf.text = f"Overall rating: {report.overall_rating.upper()}"
        tf.paragraphs[0].font.size = Pt(24)
        tf.paragraphs[0].font.bold = True

        body_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.8), prs.slide_width - Inches(1.4), Inches(3.5))
        body_tf = body_box.text_frame
        body_tf.word_wrap = True
        body_tf.text = report.overall_conclusion or "(pending)"
        for p in body_tf.paragraphs:
            p.font.size = Pt(16)

    def _signoff_slide(self, prs: Presentation, report: ValidationReport) -> None:
        slide = self._blank_slide(prs, "Sign-off")
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), prs.slide_width - Inches(1.4), Inches(1.6))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = f"Prepared by: {report.prepared_by}"
        if report.preparer:
            prep = tf.add_paragraph()
            prep.text = f"Case file prepared by: {report.preparer} (may not sign off this report)"
        req = tf.add_paragraph()
        req.text = f"Required sign-offs: {report.required_signoffs}"
        if report.signoffs:
            for signoff in report.signoffs:
                role_suffix = f" ({signoff.role})" if signoff.role else ""
                p = tf.add_paragraph()
                p.text = f"Signed off by {signoff.by}{role_suffix} at {signoff.at}"
        else:
            p = tf.add_paragraph()
            p.text = "PENDING -- human validator sign-off required"
        for p in tf.paragraphs:
            p.font.size = Pt(16)

        disclaimer_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.3), prs.slide_width - Inches(1.4), Inches(3.2))
        d_tf = disclaimer_box.text_frame
        d_tf.word_wrap = True
        d_tf.text = report.disclaimer
        d_tf.paragraphs[0].font.size = Pt(13)
        d_tf.paragraphs[0].font.italic = True
